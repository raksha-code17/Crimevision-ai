"""PDF (ReportLab) and optional DOCX export of the complete case report."""
import io
from datetime import datetime
from xml.sax.saxutils import escape

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from flask import current_app  # noqa: E402
from reportlab.lib import colors  # noqa: E402
from reportlab.lib.pagesizes import A4  # noqa: E402
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet  # noqa: E402
from reportlab.lib.units import cm  # noqa: E402
from PIL import Image as PILImage  # noqa: E402
from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

from .fusion_service import fmt_ts  # noqa: E402

INK = colors.HexColor("#17212B")
MARKER = colors.HexColor("#F2B705")


def _activity_chart(session):
    frames = [f for v in session["results"].get("video", []) for f in v["timeline"]]
    if len(frames) < 2:
        return None
    t = [f["t"] for f in frames]
    fig, ax = plt.subplots(figsize=(7, 2.2), dpi=150)
    ax.fill_between(t, [f["motion"] * 100 for f in frames], color="#2F5D8A", alpha=0.35, label="Motion (%)")
    ax.step(t, [f["persons"] for f in frames], color="#17212B", where="post", label="People")
    for f in frames:
        if f["weapons"]:
            ax.axvline(f["t"], color="#C8412B", lw=1.5)
    ax.set_xlabel("Seconds")
    ax.legend(loc="upper right", fontsize=7)
    ax.spines[["top", "right"]].set_visible(False)
    buf = io.BytesIO()
    fig.tight_layout()
    fig.savefig(buf, format="png")
    plt.close(fig)
    buf.seek(0)
    return buf


def _picture(session, rel_path, max_w=8.5 * cm, max_h=6.5 * cm):
    """Embed a saved annotated image, scaled to fit. Returns None if the file is missing."""
    path = current_app.config["UPLOAD_FOLDER"] / session["id"] / rel_path
    if not rel_path or not path.exists():
        return None
    with PILImage.open(path) as im:
        w, h = im.size
    scale = min(max_w / w, max_h / h)
    return Image(str(path), width=w * scale, height=h * scale)


def _tag_line(tags):
    return ", ".join(f"{t['label']} ({int(t['score'] * 100)}%)" for t in tags) or "none"


def _styles():
    ss = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("t", parent=ss["Title"], textColor=INK, alignment=0, fontSize=22),
        "h": ParagraphStyle("h", parent=ss["Heading2"], textColor=INK, spaceBefore=14, spaceAfter=6),
        "b": ParagraphStyle("b", parent=ss["BodyText"], leading=14),
        "small": ParagraphStyle("s", parent=ss["BodyText"], fontSize=8, textColor=colors.grey),
    }


def _table(rows, widths, header=True):
    t = Table(rows, colWidths=widths, repeatRows=1 if header else 0)
    style = [("FONTSIZE", (0, 0), (-1, -1), 8), ("VALIGN", (0, 0), (-1, -1), "TOP"),
             ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C9D1D6")),
             ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F5F6")])]
    if header:
        style += [("BACKGROUND", (0, 0), (-1, 0), MARKER), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold")]
    t.setStyle(TableStyle(style))
    return t


def generate_pdf(session: dict):
    fusion = session.get("fusion")
    if not fusion:
        raise ValueError("Run the analysis before generating a report.")
    out_dir = current_app.config["REPORT_FOLDER"]
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"report_{session['id']}.pdf"
    st = _styles()
    P = lambda text, style="b": Paragraph(escape(str(text)), st[style])  # noqa: E731
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
                            title=f"CrimeVision AI report {session['id']}", author="CrimeVision AI")
    story = [P("CrimeVision AI", "title"), P(f"Forensic analysis report · {session['title']}"),
             P(f"Case {session['id']} · generated {datetime.now():%Y-%m-%d %H:%M}", "small"), Spacer(1, 10)]

    crime, cons = fusion["crime_type"], fusion["consistency"]
    story += [P("Incident classification", "h"), P(f"{crime['type']} — confidence {int(crime['confidence'] * 100)}%"), P(crime["rationale"]),
              P(f"Cross-modal consistency: {cons['score'] if cons['score'] is not None else 'n/a'}{'%' if cons['score'] is not None else ''}. {cons['summary']}")]

    story.append(P("Key facts (who, where, when)", "h"))
    facts = fusion["facts"]
    story.append(_table([["Who", "; ".join(facts["who"]) or "—"], ["Where", "; ".join(facts["where"]) or "—"], ["When", "; ".join(facts["when"]) or "—"]], [3 * cm, 14 * cm], header=False))

    story.append(P("Evidence files", "h"))
    rows = [["ID", "Type", "File"]] + [[f["id"], f["kind"], f["name"]] for f in session["files"]]
    story.append(_table(rows, [2 * cm, 3 * cm, 12 * cm]))

    if session["results"].get("images"):
        story.append(P("Image analysis", "h"))
        rows = [["ID", "Scene", "People", "Weapons", "Summary"]] + [
            [Paragraph(escape(i["evidence_id"]), st["small"]), Paragraph(escape(i["scene"]["label"]), st["small"]), i["counts"]["persons"], (i["counts"]["weapons"] if i["counts"]["weapons"] else ("possible: " + ", ".join(i["weapon_hints"]) if i.get("weapon_hints") else ("0" if i.get("weapons_checked", True) else "not checked"))), Paragraph(escape(i["summary"]), st["small"])]
            for i in session["results"]["images"]]
        story.append(_table(rows, [1.5 * cm, 3.5 * cm, 1.5 * cm, 1.7 * cm, 8.8 * cm]))

    for img in session["results"].get("images", []):
        pic = _picture(session, img.get("annotated"))
        block = [P(f"{img['evidence_id']} · {img['file']}"), P(img["summary"], "small"), P("Classifier indicators (unverified): " + _tag_line(img.get("tags", [])), "small")]
        if (img.get("vlm") or {}).get("evidence_items"):
            block.append(P("Items noted by vision model: " + "; ".join(img["vlm"]["evidence_items"]), "small"))
        story.append(KeepTogether(([pic] if pic else []) + block + [Spacer(1, 8)]))

    for vid in session["results"].get("video", []):
        story.append(P(f"Video {vid['evidence_id']} · {vid['file']}", "h"))
        if vid.get("narrative"):
            story.append(P(vid["narrative"]))
        if vid.get("actions"):
            story.append(P("Actions noted: " + "; ".join(vid["actions"]), "small"))
        cells = []
        for kf in vid.get("keyframes", []):
            pic = _picture(session, kf["file"], max_w=5.2 * cm, max_h=4 * cm)
            caption = Paragraph(escape(f"{fmt_ts(kf['t'])} {kf['description'] or _tag_line(kf['tags'])}")[:200], st["small"])
            cells.append([pic, caption] if pic else [caption])
        for i in range(0, len(cells), 3):
            row = cells[i : i + 3]
            story.append(Table([row + [""] * (3 - len(row))], colWidths=[5.6 * cm] * 3, style=[("VALIGN", (0, 0), (-1, -1), "TOP")]))

    chart = _activity_chart(session)
    if chart:
        story += [P("Video activity", "h"), Image(chart, width=16 * cm, height=5 * cm)]

    story.append(P("Chronological timeline", "h"))
    rows = [["Time", "Source", "Event"]] + [[fmt_ts(i["t"]), i["source"], Paragraph(escape(i["text"]), st["small"])] for i in fusion["timeline"][:120]]
    story.append(_table(rows, [2 * cm, 2.2 * cm, 12.8 * cm]))
    if len(fusion["timeline"]) > 120:
        story.append(P(f"… {len(fusion['timeline']) - 120} further entries omitted.", "small"))

    story.append(P("Statement verification", "h"))
    if cons["claims"]:
        rows = [["Time", "Check", "Status", "Detail"]] + [[fmt_ts(c["t"]), c["type"], c["status"], Paragraph(escape(c["evidence"]), st["small"])] for c in cons["claims"]]
        story.append(_table(rows, [2 * cm, 2 * cm, 2.6 * cm, 10.4 * cm]))
    else:
        story.append(P("No verifiable statements found."))

    story.append(P("AI scenario hypotheses", "h"))
    for n, h in enumerate(fusion["hypotheses"]["items"], 1):
        story += [P(f"{n}. {h['title']} ({int(h['confidence'] * 100)}%)"), P(h["narrative"])]
        if h["supporting_evidence"]:
            story.append(P("Supporting: " + "; ".join(h["supporting_evidence"]), "small"))
        if h["gaps"]:
            story.append(P("Gaps: " + "; ".join(h["gaps"]), "small"))
        story.append(Spacer(1, 6))

    if session["chat"]:
        story.append(P("Investigator Q&A", "h"))
        for m in session["chat"]:
            story.append(P(("Q: " if m["role"] == "user" else "A: ") + m["content"]))

    story.append(P("Analysis methods used", "h"))
    rows = [["ID", "Method"]]
    for i in session["results"].get("images", []) + session["results"].get("video", []):
        e = i["engines"]
        rows.append([i["evidence_id"], f"Object detector: {e['detector']}; CLIP indicators: {'yes' if e['clip'] else 'no'}; vision model: {'yes' if e['vision_model'] else 'no'}"])
    rows += [[a["evidence_id"], f"Speech/transcript: {a['engine']}; facts: {a['facts'].get('engine', 'n/a')}"] for a in session["results"].get("audio", [])]
    story.append(_table(rows, [2 * cm, 15 * cm]))
    story.append(P("Classifier indicators and vision-model descriptions are possibilities, not confirmed findings. Fewer methods means less coverage.", "small"))

    story += [Spacer(1, 12), P("AI-generated analysis for investigative support only. Verify all findings against original evidence before relying on them.", "small")]
    doc.build(story)
    return path


def generate_docx(session: dict):
    from docx import Document

    fusion = session.get("fusion")
    if not fusion:
        raise ValueError("Run the analysis before generating a report.")
    out_dir = current_app.config["REPORT_FOLDER"]
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"report_{session['id']}.docx"
    d = Document()
    d.add_heading("CrimeVision AI — Forensic report", 0)
    d.add_paragraph(f"{session['title']} · case {session['id']}")
    d.add_heading("Classification", 1)
    d.add_paragraph(f"{fusion['crime_type']['type']} ({int(fusion['crime_type']['confidence'] * 100)}%): {fusion['crime_type']['rationale']}")
    d.add_paragraph(fusion["consistency"]["summary"])
    d.add_heading("Key facts", 1)
    for k, v in fusion["facts"].items():
        d.add_paragraph(f"{k.capitalize()}: {'; '.join(v) or '—'}")
    d.add_heading("Timeline", 1)
    for i in fusion["timeline"]:
        d.add_paragraph(f"[{fmt_ts(i['t'])}] ({i['source']}) {i['text']}")
    d.add_heading("Hypotheses", 1)
    for h in fusion["hypotheses"]["items"]:
        d.add_paragraph(f"{h['title']} ({int(h['confidence'] * 100)}%): {h['narrative']}")
    d.save(str(path))
    return path
